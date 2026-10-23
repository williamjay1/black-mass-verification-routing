"""Annual, auditable NMC-reference supply for an EU27 engineering design.

UMP's Li-rechargeable layer is cell-based; PackXEV excludes cells/modules.
Collection, NMC mass share, and transferring OPARA case yields are declared
scenarios, not measurements of European black-mass supply.
"""
from __future__ import annotations
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parent
FLOW = ROOT / "datasets/ump_bev_flows_20260930_v1/eu27_bev_lithium_battery_waste_2025_2050.csv"
POINTS = Path(r"E:\AcademicData\topic34\raw\gisco\NUTS_LB_2024_4326_LEVL_0_2026-09-26.geojson")
HUBS = PROJECT / "datasets/upgrade/eu27_candidate_hub_nodes_2026-09-26.csv"
YIELD = Path(r"F:\AcademicData\topic34\raw\OPARA-276_20260928\verified_downloads_v2\2_-_Black_mass_yield_and_relation_black_mass_share_to_coating_share.xlsx")
OUT = ROOT / "datasets/futuram_network_inputs_20260930_v1"
ISO = dict(zip(
    ["AUT","BEL","BGR","HRV","CYP","CZE","DNK","EST","FIN","FRA","DEU","GRC","HUN","IRL","ITA","LVA","LTU","LUX","MLT","NLD","POL","PRT","ROU","SVK","SVN","ESP","SWE"],
    ["AT","BE","BG","HR","CY","CZ","DK","EE","FI","FR","DE","EL","HU","IE","IT","LV","LT","LU","MT","NL","PL","PT","RO","SK","SI","ES","SE"]))

def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    book = openpyxl.load_workbook(YIELD, data_only=True, read_only=True)
    sheet = book["Black mass yield"]
    observed = pd.DataFrame([{"cell_type":sheet.cell(16,c).value,
                              "black_mass_yield_fraction":float(sheet.cell(17,c).value)/100,
                              "source_cell":f"Black mass yield!{openpyxl.utils.get_column_letter(c)}17"}
                            for c in range(3,13)])
    observed.to_csv(OUT / "reported_cell_to_black_mass_yield_cases.csv",index=False)
    yield_refs = {"minimum_reported_case":float(observed.black_mass_yield_fraction.min()),
                  "unweighted_case_median":float(observed.black_mass_yield_fraction.median()),
                  "maximum_reported_case":float(observed.black_mass_yield_fraction.max())}
    flow = pd.read_csv(FLOW)
    flow["country_code"] = flow.country_iso3.map(ISO)
    assert flow.country_code.notna().all() and flow.country_code.nunique() == 27
    centres = []
    for feature in json.loads(POINTS.read_text(encoding="utf-8"))["features"]:
        prop = feature["properties"]
        code = prop["CNTR_CODE"]
        if code not in set(ISO.values()):
            continue
        assert prop["LEVL_CODE"] == 0 and feature["geometry"]["type"] == "Point"
        lon, lat = feature["geometry"]["coordinates"]
        centres.append({"origin_id":code,"country_code":code,
                        "lat":float(lat),"lon":float(lon),
                        "origin_name":prop["NAME_ENGL"],
                        "spatial_representation":"GISCO 2024 NUTS-0 official label point; national geometric proxy, not collection location"})
    centres = pd.DataFrame(centres)
    assert len(centres) == 27 and set(centres.country_code) == set(ISO.values())
    centres.to_csv(OUT / "national_representative_points.csv",index=False)
    hub = pd.read_csv(HUBS)
    hub = hub.loc[hub.base_candidate.astype(str).str.lower().eq("true")].copy()
    hub.to_csv(OUT / "candidate_hubs.csv",index=False)
    distances = []
    for r in centres.itertuples():
        for h in hub.itertuples():
            phi1,phi2 = np.radians([r.lat,h.lat])
            a = np.sin((phi2-phi1)/2)**2+np.cos(phi1)*np.cos(phi2)*np.sin(np.radians(h.lon-r.lon)/2)**2
            distances.append({"origin_id":r.origin_id,"hub_id":h.hub_id,
                              "distance_km":float(6371.0088*2*np.arcsin(np.sqrt(min(1.,a))))})
    pd.DataFrame(distances).to_csv(OUT / "national_origin_hub_geodesic_matrix.csv",index=False)
    summaries = []
    for year in (2025,2030,2035,2040,2050):
        national = flow.loc[flow.year.eq(year)].merge(centres,on="country_code",validate="one_to_one")
        national["assumed_nmc_cell_mass_share"] = 0.5
        national["assumed_collection_fraction"] = 1.0
        national["reported_case_yield_reference"] = yield_refs["unweighted_case_median"]
        national["reference_black_mass_tpy"] = national.bev_lithium_battery_waste_t*0.5*yield_refs["unweighted_case_median"]
        national["flow_basis"] = "modelled BEV Li-cell waste * assumed NMC mass share * case-reference BM yield; 100% collection potential"
        national.to_csv(OUT / f"national_origins_{year}.csv",index=False)
        total = float(national.bev_lithium_battery_waste_t.sum())
        for name,bm_yield in yield_refs.items():
            for share in (0.2,0.5,0.8):
                for collection in (0.5,0.75,1.0):
                    summaries.append({"year":year,"li_cell_waste_t":total,"yield_reference":name,
                                      "yield_fraction":bm_yield,"assumed_nmc_mass_share":share,
                                      "assumed_collection_fraction":collection,
                                      "nmc_reference_black_mass_t":total*share*collection*bm_yield})
        assert np.isclose(national.reference_black_mass_tpy.sum(),total*0.5*yield_refs["unweighted_case_median"])
    pd.DataFrame(summaries).to_csv(OUT / "supply_conversion_scenario_envelope.csv",index=False)
    manifest = {"generated_utc":datetime.now(timezone.utc).isoformat(),
                "inputs":[{"path":str(p),"sha256":sha(p)} for p in (FLOW,POINTS,HUBS,YIELD)],
                "script_sha256":sha(Path(__file__)),
                "official_definition":"PackXEV excludes cells and modules; composition model is cell-level (FutuRaM battery supporting document p.6,8)",
                "base_conversion":{"nmc_mass_share":0.5,"collection_fraction":1.0,"bm_yield":yield_refs["unweighted_case_median"]},
                "yield_support":yield_refs,
                "yield_boundary":"Source-specific BM output <1000 micrometres after drying per original cell mass; case envelope, not EU/industrial mean",
                "chemistry_boundary":"0.5 is a declared mass-share reference; FutuRaM assumes equal LFP/NMC shares, whose exact mass basis is not verified",
                "geography":"27 national flows at official GISCO 2024 NUTS-0 label points; inventory city hubs; geodesic lower-bound geometry; not road logistics",
                "annual_snapshots":[2025,2030,2035,2040,2050],
                "national_origin_count":len(centres),"base_candidate_city_count":len(hub),
                "claims_excluded":["observed waste generation","collected BM tonnage","EU chemistry prevalence","industrial BM yield","permitted route capacity","real investment feasibility"]}
    (OUT / "input_manifest.json").write_text(json.dumps(manifest,indent=2),encoding="utf-8")
    print(json.dumps(manifest,indent=2),flush=True)

if __name__ == "__main__":
    main()
