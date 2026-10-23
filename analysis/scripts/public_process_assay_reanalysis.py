"""Reproduce public-data black-mass process/assay decision audits.

Reads immutable source workbooks in F:\\AcademicData and writes all derived
files to D:\\MLWork. No European prevalence, independent-lot inference, or route
compatibility label is inferred from the source experiments.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import statistics
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl


RAW_TUD = Path(r"F:\AcademicData\topic34\raw\opara_769_kaas_2025_20260930")
RAW_ANTICOI = Path(
    r"F:\AcademicData\topic34\raw\anticoi_2025_jmcwm_10.1007_s10163-025-02159-x_20260930"
)
OUT = Path(
    r"D:\MLWork\topic34\manuscript_quality_v2\results\public_process_assay_audit_2026-09-30"
)
SETTINGS = [
    "10-1500", "10-1750", "10-2000", "20-1250", "20-1500", "20-1750",
    "20-2000", "20-2250", "30-1500", "30-1750", "30-2000",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest().upper()


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError(f"Refusing to write empty table: {path}")
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def nfloat(value: Any) -> float | None:
    try:
        out = float(value)
        return out if math.isfinite(out) else None
    except (TypeError, ValueError):
        return None


def find_setting_header(ws) -> tuple[int, dict[str, int]]:
    for r, row in enumerate(ws.iter_rows(values_only=True)):
        cols: dict[str, int] = {}
        for c, value in enumerate(row):
            key = str(value).strip()
            # Some worksheets place the same 11-setting header beside the
            # mean, dMin and dMax blocks. Retain the first (mean) block.
            if key in SETTINGS and key not in cols:
                cols[key] = c
        if len(cols) == len(SETTINGS):
            return r, cols
    raise ValueError(f"Could not find all 11 settings in worksheet {ws.title}")


def extract_quality_and_recovery() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    composition_path = RAW_TUD / "5_Black_mass_composition.xlsx"
    recovery_path = RAW_TUD / "6_Element_recovery.xlsx"
    classifier_path = RAW_TUD / "4_Recovery_product_composition.xlsx"

    composition_wb = openpyxl.load_workbook(composition_path, data_only=True, read_only=True)
    ws_comp = composition_wb["BM2"]
    comp_header_row, comp_cols = find_setting_header(ws_comp)
    comp_rows: dict[str, dict[str, float]] = {}
    for row in ws_comp.iter_rows(min_row=comp_header_row + 1, values_only=True):
        if any(isinstance(v, str) and "error bars" in v.lower() for v in row):
            break
        label = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        if label in {"Al", "Co", "Cu", "Fe", "Li", "Mn", "Ni", "P", "NMC", "PF6", "Other"}:
            vals = {setting: nfloat(row[col]) for setting, col in comp_cols.items()}
            comp_rows[label] = {k: v for k, v in vals.items() if v is not None}

    recovery_wb = openpyxl.load_workbook(recovery_path, data_only=True, read_only=True)
    ws_rec = recovery_wb["BM2"]
    rec_header_row, rec_cols = find_setting_header(ws_rec)
    rec_rows: dict[str, dict[str, float]] = {}
    rec_names = {"Cu": "cu_recovery_pct", "Li": "li_recovery_pct", "Ni": "ni_recovery_pct",
                 "Co": "co_recovery_pct", "Mn": "mn_recovery_pct", "P": "p_recovery_pct"}
    for row in ws_rec.iter_rows(min_row=rec_header_row + 1, values_only=True):
        if any(isinstance(v, str) and "error bars" in v.lower() for v in row):
            break
        label = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        if label in rec_names:
            rec_rows[rec_names[label]] = {
                setting: float(row[col]) for setting, col in rec_cols.items()
                if nfloat(row[col]) is not None
            }

    classifier_wb = openpyxl.load_workbook(classifier_path, data_only=True, read_only=True)
    ws_class = classifier_wb["Fuerstenau 2_anode,cathode"]
    blocks: list[tuple[int, dict[str, int], str]] = []
    for r, row in enumerate(ws_class.iter_rows(values_only=True)):
        cols: dict[str, int] = {}
        for c, value in enumerate(row):
            key = str(value).strip()
            if key in SETTINGS and key not in cols:
                cols[key] = c
        if len(cols) == len(SETTINGS):
            cut_label = "unknown"
            for rr in range(max(0, r - 5), r):
                prior_row = next(ws_class.iter_rows(min_row=rr + 1, max_row=rr + 1, values_only=True))
                joined = " ".join(str(val).strip() for val in prior_row if isinstance(val, str))
                if "cut size" in joined.lower():
                    cut_label = joined
            blocks.append((r, cols, cut_label))
    if len(blocks) != 2:
        raise ValueError(f"Expected two classifier blocks, found {len(blocks)}")

    class_rows: dict[tuple[str, str], float] = {}
    label_map = {
        "Cathode into cathode product": "cathode_recovery_pct",
        "Cathode into anode product": "cathode_carryover_to_anode_pct",
        "Anode into Anode product": "anode_recovery_pct",
        "Anode into cathode product": "anode_carryover_to_cathode_pct",
        "Conc. anode in anode product": "anode_concentration_pct",
    }
    all_rows = list(ws_class.iter_rows(values_only=True))
    for block_i, (header_r, cols, cut_label) in enumerate(blocks):
        next_header = blocks[block_i + 1][0] if block_i + 1 < len(blocks) else len(all_rows)
        cut = "3.2_m_per_s" if "3.2" in cut_label else "4.2_m_per_s" if "4.2" in cut_label else cut_label
        for row in all_rows[header_r + 1:next_header]:
            text = " | ".join(str(v).strip() for v in row if isinstance(v, str))
            for source_label, output_name in label_map.items():
                if source_label in text:
                    for setting, col in cols.items():
                        val = nfloat(row[col])
                        if val is not None:
                            class_rows[(cut, output_name + "|" + setting)] = val

    merged: list[dict[str, Any]] = []
    for setting in SETTINGS:
        item: dict[str, Any] = {"setting": setting}
        for element in ("Al", "Cu", "Fe", "Li", "Ni", "Co", "Mn", "P", "NMC", "PF6", "Other"):
            item[f"{element.lower()}_wt_pct"] = comp_rows.get(element, {}).get(setting)
        item["cu_al_wt_pct"] = item["cu_wt_pct"] + item["al_wt_pct"]
        item["cu_al_fe_p_wt_pct"] = sum(item[f"{el.lower()}_wt_pct"] for el in ("Cu", "Al", "Fe", "P"))
        for element, output_name in (("Cu", "cu_recovery_pct"), ("Li", "li_recovery_pct"),
                                    ("Ni", "ni_recovery_pct"), ("Co", "co_recovery_pct"),
                                    ("Mn", "mn_recovery_pct"), ("P", "p_recovery_pct")):
            item[output_name] = rec_rows.get(output_name, {}).get(setting)
        for cut in ("3.2_m_per_s", "4.2_m_per_s"):
            for metric in ("cathode_recovery_pct", "cathode_carryover_to_anode_pct",
                           "anode_recovery_pct", "anode_carryover_to_cathode_pct",
                           "anode_concentration_pct"):
                item[f"{cut}_{metric}"] = class_rows.get((cut, metric + "|" + setting))
        item["passes_illustrative_4wtpct_cu_al"] = item["cu_al_wt_pct"] <= 4.0
        merged.append(item)

    # These pairs are complementary source-reported component recoveries.
    # The identity is checked setting-by-setting, not used to infer new data.
    for row in merged:
        for cut in ("3.2_m_per_s", "4.2_m_per_s"):
            cathode_sum = row[f"{cut}_cathode_recovery_pct"] + row[f"{cut}_cathode_carryover_to_anode_pct"]
            anode_sum = row[f"{cut}_anode_recovery_pct"] + row[f"{cut}_anode_carryover_to_cathode_pct"]
            if abs(cathode_sum - 100.0) > 0.02 or abs(anode_sum - 100.0) > 0.02:
                raise ValueError(f"Classifier material balance failed for {row['setting']} at {cut}")

    metadata = {
        "source_doi": "10.1016/j.wasman.2025.114912",
        "repository_doi": "10.25532/OPARA-769",
        "repository_url": "https://opara.zih.tu-dresden.de/items/49ef6eba-8b70-4c54-98ab-49e9fdaf3909/full",
        "license": "CC BY 4.0",
        "source_date": "2025-03-12",
        "worksheet_sources": {
            "BM2 composition": {"file": composition_path.name, "sheet": "BM2", "header_row_1based": comp_header_row + 1},
            "BM2 element recovery": {"file": recovery_path.name, "sheet": "BM2", "header_row_1based": rec_header_row + 1},
            "air classification": {"file": classifier_path.name, "sheet": ws_class.title,
                                    "cut_labels": [b[2] for b in blocks], "header_rows_1based": [b[0] + 1 for b in blocks]},
        },
        "raw_sha256": {p.name: sha256(p) for p in (composition_path, recovery_path, classifier_path)},
        "n_settings": len(merged),
        "design_unit_warning": "Eleven settings from the source experimental grid, not eleven independent batteries or European lots.",
    }
    return merged, metadata


def assay_analysis() -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    # Article-reported values, Table 3; percent as printed. These are three
    # size-separated streams from one Spanish micromobility feed, not replicates.
    # "N.D." for XRF lithium is retained as missing/non-detected, never zero.
    rows = [
        ("A5", "Al", 4.7, 6.7), ("A5", "Co", 3.4, 5.6),
        ("A5", "Cu", 1.1, 1.7), ("A5", "Li", 2.7, None),
        ("A5", "Mn", 4.0, 4.9), ("A5", "Ni", 26.4, 38.8),
        ("A6", "Al", 10.0, 16.0), ("A6", "Co", 2.8, 2.2),
        ("A6", "Cu", 15.1, 16.1), ("A6", "Li", 2.4, None),
        ("A6", "Mn", 3.6, 2.6), ("A6", "Ni", 17.8, 17.5),
        ("A7", "Al", 6.4, 12.4), ("A7", "Co", 0.3, 0.25),
        ("A7", "Cu", 44.7, 35.7), ("A7", "Li", 0.2, None),
        ("A7", "Mn", 0.2, 0.2), ("A7", "Ni", 2.2, 2.5),
    ]
    elemental = [
        {"stream": s, "element": e, "aas_reported_pct": a, "xrf_reported_pct": x,
         "xrf_status": "not_detected_or_not_reported_as_numeric" if x is None else "numeric",
         "difference_xrf_minus_aas_percentage_points": None if x is None else x - a,
         "relative_difference_vs_aas_pct": None if x is None or a == 0 else (x - a) / a * 100.0}
        for s, e, a, x in rows
    ]
    sums = []
    for stream in ("A5", "A6", "A7"):
        vals = {e: (a, x) for s, e, a, x in rows if s == stream}
        aas_sum = sum(vals[e][0] for e in ("Cu", "Al"))
        xrf_sum = sum(vals[e][1] for e in ("Cu", "Al"))
        sums.append({"stream": stream, "aas_cu_al_pct": aas_sum, "xrf_cu_al_pct": xrf_sum,
                     "xrf_minus_aas_percentage_points": xrf_sum - aas_sum,
                     "classification_discordance_threshold_low_inclusive_pct": min(aas_sum, xrf_sum),
                     "classification_discordance_threshold_high_exclusive_pct": max(aas_sum, xrf_sum)})
    summary = {
        "source": "Anticoi et al. (2025), Table 3, DOI 10.1007/s10163-025-02159-x",
        "article_url": "https://link.springer.com/article/10.1007/s10163-025-02159-x",
        "open_full_text_url": "https://upcommons.upc.edu/server/api/core/bitstreams/16f53294-beff-4f46-a069-7e94ec2312e0/content",
        "raw_supplement": "10163_2025_2159_MOESM1_ESM.xlsx",
        "supplement_sha256": sha256(RAW_ANTICOI / "10163_2025_2159_MOESM1_ESM.xlsx"),
        "n_paired_streams": 3,
        "n_numeric_element_method_pairs": sum(x is not None for _, _, _, x in rows),
        "element_descriptive_absolute_differences_percentage_points": {
            element: {
                "n_numeric_pairs": sum(1 for s, e, a, x in rows if e == element and x is not None),
                "median_absolute_difference": statistics.median(
                    abs(x - a) for s, e, a, x in rows if e == element and x is not None
                ) if any(e == element and x is not None for s, e, a, x in rows) else None,
            }
            for element in ("Al", "Co", "Cu", "Li", "Mn", "Ni")
        },
        "xrf_lithium_status": "All three Table 3 XRF entries are reported N.D.; these are not interpreted as zero or as a detection-limit estimate.",
        "threshold_rule": "pass when Cu+Al <= threshold; disagreement interval is [min(AAS,XRF), max(AAS,XRF)).",
        "4wtpct_cutoff_observed_flips": 0,
        "interpretation_limit": "Arithmetic method disagreement only; AAS is not a validated route-compatibility label; no sensitivity/specificity or European prevalence is estimated.",
    }
    return elemental, sums, summary


def synthetic_xrf_reanalysis() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Compare four-metal proportions in eight designed mixes.

    The response is deliberately a within-four-metal composition, not a
    reconstruction of Anticoi et al.'s LOI/carbon-corrected whole-sample
    estimate. XRF oxide percentages are converted to elemental equivalents,
    then both measured and recipe elemental masses are normalized over Cu,
    Mn, Co, and Ni. This removes the unknown common denominator while keeping
    the result matrix-specific and non-inferential.
    """
    path = RAW_ANTICOI / "10163_2025_2159_MOESM1_ESM.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["XRF Results"]
    all_rows = list(ws.iter_rows(values_only=True))
    labels = [str(v).strip() if v is not None else "" for v in all_rows[0]]
    sample_cols = {f"R{i}": labels.index(f"R{i}") for i in range(1, 9)}
    oxide_labels = {str(row[4]).strip().rstrip() for row in all_rows if len(row) > 4 and row[4] is not None}
    oxide_row: dict[str, tuple[Any, ...]] = {}
    for row in all_rows:
        if len(row) > 4 and row[4] is not None:
            oxide_row[str(row[4]).strip().rstrip()] = row
    source_oxide = {"Mn": "MnO", "Co": "Co3O4", "Ni": "NiO", "Cu": "CuO"}
    xrf_to_metal = {
        "Mn": 54.938044 / (54.938044 + 15.999),
        "Co": (3 * 58.933194) / (3 * 58.933194 + 4 * 15.999),
        "Ni": 58.6934 / (58.6934 + 15.999),
        "Cu": 63.546 / (63.546 + 15.999),
    }
    # Grams in the source article's Table 2 (R1--R8); recipe rows are the
    # prepared Cu, MnO2, NiO, LiCoO2 and activated-carbon components.
    recipe = {
        "Cu": [2.011, 1.007, 4.005, 2.050, 0.506, 1.030, 0.250, 0.510],
        "MnO2": [1.999, 1.005, 2.010, 1.018, 0.250, 1.010, 0.252, 0.500],
        "NiO": [6.004, 3.010, 4.011, 1.997, 0.508, 1.001, 0.753, 1.520],
        "LiCoO2": [2.005, 1.116, 2.004, 1.030, 0.253, 0.520, 0.249, 0.506],
    }
    recipe_to_metal = {
        "Cu": 1.0,
        "MnO2": 54.938044 / (54.938044 + 2 * 15.999),
        "NiO": 58.6934 / (58.6934 + 15.999),
        "LiCoO2": 58.933194 / (6.94 + 58.933194 + 2 * 15.999),
    }
    output: list[dict[str, Any]] = []
    for idx in range(8):
        sample = f"R{idx + 1}"
        xrf_equiv: dict[str, float] = {}
        recipe_metal: dict[str, float] = {}
        for element, oxide in source_oxide.items():
            value = nfloat(oxide_row[oxide][sample_cols[sample]])
            if value is None:
                raise ValueError(f"Missing {oxide} XRF value for {sample}")
            xrf_equiv[element] = value * xrf_to_metal[element]
        recipe_metal = {
            "Cu": recipe["Cu"][idx] * recipe_to_metal["Cu"],
            "Mn": recipe["MnO2"][idx] * recipe_to_metal["MnO2"],
            "Ni": recipe["NiO"][idx] * recipe_to_metal["NiO"],
            "Co": recipe["LiCoO2"][idx] * recipe_to_metal["LiCoO2"],
        }
        xrf_total = sum(xrf_equiv.values())
        recipe_total = sum(recipe_metal.values())
        for element in ("Cu", "Mn", "Co", "Ni"):
            xrf_share = 100 * xrf_equiv[element] / xrf_total
            recipe_share = 100 * recipe_metal[element] / recipe_total
            output.append({
                "sample": sample,
                "element": element,
                "xrf_oxide_label": source_oxide[element],
                "xrf_reported_oxide_wt_pct": nfloat(oxide_row[source_oxide[element]][sample_cols[sample]]),
                "xrf_element_share_within_four_metals_pct": xrf_share,
                "recipe_element_share_within_four_metals_pct": recipe_share,
                "difference_percentage_points": xrf_share - recipe_share,
                "absolute_difference_percentage_points": abs(xrf_share - recipe_share),
            })
    abs_errors = [row["absolute_difference_percentage_points"] for row in output]
    summary = {
        "source_article": "Anticoi et al. (2025), Table 2 and Supplementary Information, XRF Results sheet",
        "source_doi": "10.1007/s10163-025-02159-x",
        "raw_supplement_sha256": sha256(path),
        "n_designed_mixtures": 8,
        "n_element_proportions": len(output),
        "normalization": "Oxide-normalized XRF values were converted to Cu, Mn, Co and Ni elemental-equivalent masses using stoichiometric factors, then normalized over those four metals. The known recipe was independently converted to the same four-metal basis.",
        "not_reconstructed": "This is not a reconstruction of the authors' whole-sample LOI/activated-carbon correction or Figure 7, and not a field-feed accuracy estimate.",
        "mean_absolute_difference_percentage_points": sum(abs_errors) / len(abs_errors),
        "maximum_absolute_difference_percentage_points": max(abs_errors),
        "scope_limit": "Eight deliberately prepared pseudocomponent mixtures; the source does not establish independent replicates or coverage of commercial lot matrices. No direct-process compatibility labels are present.",
    }
    return output, summary


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    process_rows, process_metadata = extract_quality_and_recovery()
    assay_elements, assay_sums, assay_metadata = assay_analysis()
    synthetic_rows, synthetic_metadata = synthetic_xrf_reanalysis()
    write_csv(OUT / "kaas_2025_process_quality_recovery.csv", process_rows)
    write_csv(OUT / "anticoi_2025_xrf_aas_element_differences.csv", assay_elements)
    write_csv(OUT / "anticoi_2025_cu_al_threshold_disagreement_intervals.csv", assay_sums)
    write_csv(OUT / "anticoi_2025_synthetic_four_metal_proportions.csv", synthetic_rows)

    cu_al = [row["cu_al_wt_pct"] for row in process_rows]
    cu_only = [row["cu_wt_pct"] for row in process_rows]
    nmc = [row["nmc_wt_pct"] for row in process_rows]
    summary = {
        "run_date": date.today().isoformat(),
        "analysis_unit": "public source-reported setting/stream; no inferential population model",
        "kaas": {
            **process_metadata,
            "cu_wt_pct_range": [min(cu_only), max(cu_only)],
            "cu_al_wt_pct_range": [min(cu_al), max(cu_al)],
            "nmc_wt_pct_range": [min(nmc), max(nmc)],
            "count_at_or_below_illustrative_4wtpct_cu_al": sum(v <= 4.0 for v in cu_al),
            "count_at_or_below_illustrative_4wtpct_cu_al_fe_p": sum(row["cu_al_fe_p_wt_pct"] <= 4.0 for row in process_rows),
            "unweighted_setting_tallies_by_cu_al_cutoff_wt_pct": {
                str(t): sum(v <= t for v in cu_al) for t in (2, 3, 4, 5, 6, 8, 10, 15, 20)
            },
            "four_wt_pct_cutoff_status": "Analyst-selected sensitivity value, not a published process acceptance limit.",
            "fixed_10mm_speed_sweep": [row for row in process_rows if row["setting"].startswith("10-")],
            "classification_recovery_ranges": {
                cut: {
                    "cathode_recovery_pct": [min(r[f"{cut}_cathode_recovery_pct"] for r in process_rows), max(r[f"{cut}_cathode_recovery_pct"] for r in process_rows)],
                    "anode_recovery_pct": [min(r[f"{cut}_anode_recovery_pct"] for r in process_rows), max(r[f"{cut}_anode_recovery_pct"] for r in process_rows)],
                    "anode_concentration_pct": [min(r[f"{cut}_anode_concentration_pct"] for r in process_rows), max(r[f"{cut}_anode_concentration_pct"] for r in process_rows)],
                }
                for cut in ("3.2_m_per_s", "4.2_m_per_s")
            },
            "paired_cutpoint_changes_4p2_minus_3p2": {
                metric: {
                    "mean_across_11_source_settings": sum(r[f"4.2_m_per_s_{metric}"] - r[f"3.2_m_per_s_{metric}"] for r in process_rows) / len(process_rows),
                    "min_across_settings": min(r[f"4.2_m_per_s_{metric}"] - r[f"3.2_m_per_s_{metric}"] for r in process_rows),
                    "max_across_settings": max(r[f"4.2_m_per_s_{metric}"] - r[f"3.2_m_per_s_{metric}"] for r in process_rows),
                }
                for metric in ("cathode_recovery_pct", "anode_recovery_pct", "anode_concentration_pct")
            },
        },
        "anticoi": {**assay_metadata, "threshold_disagreement_intervals": assay_sums,
                     "synthetic_four_metal_proportion_audit": synthetic_metadata},
        "reproducibility": {
            "script": "public_process_assay_reanalysis.py",
            "outputs": ["kaas_2025_process_quality_recovery.csv", "anticoi_2025_xrf_aas_element_differences.csv", "anticoi_2025_cu_al_threshold_disagreement_intervals.csv", "anticoi_2025_synthetic_four_metal_proportions.csv", "audit_summary.json"],
            "no_raw_files_modified": True,
        },
    }
    (OUT / "audit_summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"output": str(OUT), "n_process_settings": len(process_rows),
                      "cu_al_range": summary["kaas"]["cu_al_wt_pct_range"],
                      "aas_xrf_intervals": assay_sums}, indent=2))


if __name__ == "__main__":
    main()
