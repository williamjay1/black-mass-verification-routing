#!/usr/bin/env python3
"""Read-only audit of the public OPARA-276 cell-type black-mass supplement.

The script extracts reported mean Al and Cu concentrations and computes an
equal-weight tally of the ten named cell types under analyst-selected Cu+Al
cutoffs. It does not estimate EU prevalence, assay accuracy, or route
compatibility. All outputs are written to D:\\MLWork; source files on F: are
opened read-only and are never modified.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import platform
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


RAW_ROOT = Path(r"F:\AcademicData\topic34\raw\OPARA-276_20260928")
DEFAULT_OUTPUT = Path(
    r"D:\MLWork\topic34\manuscript_quality_v2\results\opara276_celltype_gate_audit_2026-09-28"
)
DATASET_DOI = "10.25532/OPARA-276"
ARTICLE_DOI = "10.1016/j.nxsust.2024.100050"

SOURCES = {
    "readme": RAW_ROOT / "Readme.txt",
    "cell_composition": RAW_ROOT / "1_-_Cell_composition_and_element_content.xlsx",
    "black_mass_yield": RAW_ROOT
    / "verified_downloads_v2"
    / "2_-_Black_mass_yield_and_relation_black_mass_share_to_coating_share.xlsx",
    "particle_size": RAW_ROOT
    / "verified_downloads_v2"
    / "3_-_Particle_size_distribution_of_black_mass.xlsx",
    "black_mass_composition": RAW_ROOT / "4_-_Black_mass_composition_and_element_recoveries_v2.xlsx",
}

EXPECTED_HTML_ERROR_PAGES = [
    RAW_ROOT / "2_-_Black_mass_yield_and_relation_black_mass_share_to_coating_share.xlsx",
    RAW_ROOT / "3_-_Particle_size_distribution_of_black_mass.xlsx",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def audit_xlsx(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(path)
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        has_workbook = "xl/workbook.xml" in archive.namelist()
    return {
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "zip_test": bad_member,
        "has_xlsx_workbook_xml": has_workbook,
        "valid_xlsx": bad_member is None and has_workbook,
    }


def numeric(value: Any, *, coordinate: str) -> float:
    if not isinstance(value, (int, float)):
        raise ValueError(f"Expected numeric value at {coordinate}, got {value!r}")
    return float(value)


def extract_metrics(workbook_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Element content"]
    cell_types = [sheet.cell(row=8, column=column).value for column in range(4, 14)]
    if len(cell_types) != 10 or any(not isinstance(value, str) for value in cell_types):
        raise ValueError(f"Unexpected cell-type header in OPARA workbook: {cell_types!r}")

    row_by_element = {"Al": 10, "Cu": 12}
    means: dict[str, list[float]] = {}
    for element, row in row_by_element.items():
        means[element] = [
            numeric(sheet.cell(row=row, column=column).value,
                    coordinate=f"Element content!{row},{column}")
            for column in range(4, 14)
        ]

    return [
        {
            "cell_type": cell_types[index],
            "Al_mean_wt_pct": means["Al"][index],
            "Cu_mean_wt_pct": means["Cu"][index],
            "Cu_plus_Al_mean_wt_pct": means["Cu"][index] + means["Al"][index],
            "source_units": "wt% of reported black-mass composition",
            "source_sheet": "Element content",
            "source_cells": f"Al row 10, Cu row 12; cell-type columns D:M; header row 8",
        }
        for index in range(10)
    ]


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError(f"Refusing to write an empty CSV: {path}")
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def build_threshold_rows(metrics: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for half_units in range(4, 17):  # 2.0 to 8.0 wt%, inclusive, in 0.5 steps
        threshold = half_units / 2
        passing = [
            row["cell_type"]
            for row in metrics
            if row["Cu_plus_Al_mean_wt_pct"] <= threshold
        ]
        rows.append(
            {
                "Cu_plus_Al_cutoff_wt_pct": threshold,
                "n_cell_types_passing": len(passing),
                "n_cell_types_total": len(metrics),
                "equal_weight_type_fraction": len(passing) / len(metrics),
                "passing_cell_type_codes": ";".join(passing),
                "interpretation": "descriptive equal-weight cell-type tally; not feed prevalence",
            }
        )
    return rows


def build_figure(metrics: list[dict[str, Any]], threshold_rows: list[dict[str, Any]], output: Path) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.titlesize": 10,
            "axes.labelsize": 9,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "svg.fonttype": "none",
        }
    )
    cutoff = 4.0
    ordered = sorted(metrics, key=lambda row: row["Cu_plus_Al_mean_wt_pct"])
    labels = [row["cell_type"] for row in ordered]
    values = [row["Cu_plus_Al_mean_wt_pct"] for row in ordered]
    colors = ["#267A72" if value <= cutoff else "#B56A3B" for value in values]
    thresholds = [row["Cu_plus_Al_cutoff_wt_pct"] for row in threshold_rows]
    counts = [row["n_cell_types_passing"] for row in threshold_rows]

    fig, (left, right) = plt.subplots(
        1, 2, figsize=(10.2, 4.5), gridspec_kw={"width_ratios": [1.15, 1.0]}
    )
    y = list(range(len(ordered)))
    left.barh(y, values, color=colors, height=0.68, edgecolor="white", linewidth=0.5)
    left.axvline(cutoff, color="#202020", linewidth=1.25, linestyle=(0, (4, 3)))
    left.set_yticks(y, labels)
    left.invert_yaxis()
    left.set_xlim(0, 8.6)
    left.set_xlabel("Cu + Al in reported BM means (wt%)")
    left.set_title("Ten source-defined cell types")
    left.grid(axis="x", color="#D9DEE2", linewidth=0.6)
    left.set_axisbelow(True)
    for yi, value in zip(y, values):
        left.text(value + 0.10, yi, f"{value:.2f}", va="center", fontsize=8)

    right.step(thresholds, counts, where="post", color="#315A7D", linewidth=2)
    right.scatter(thresholds, counts, color="#315A7D", s=15, zorder=3)
    right.axvline(cutoff, color="#202020", linewidth=1.25, linestyle=(0, (4, 3)))
    right.scatter([cutoff], [6], color="#B56A3B", edgecolor="white", linewidth=0.7, s=65, zorder=4)
    right.annotate("4 wt%: 6/10 types", (cutoff, 6), xytext=(4.25, 4.1),
                   arrowprops={"arrowstyle": "-", "color": "#555555", "linewidth": 0.8},
                   fontsize=8)
    right.set_xlim(1.8, 8.2)
    right.set_ylim(-0.25, 10.8)
    right.set_xticks([2, 3, 4, 5, 6, 7, 8])
    right.set_yticks([0, 2, 4, 6, 8, 10])
    right.set_xlabel("Analyst-selected Cu + Al cutoff (wt%)")
    right.set_ylabel("Cell-type means at or below cutoff (n/10)")
    right.set_title("Cutoff sensitivity of the type tally")
    right.grid(color="#D9DEE2", linewidth=0.6)
    right.set_axisbelow(True)

    fig.suptitle("Published cell-type black-mass means do not define batch prevalence or route labels", y=1.02,
                 fontsize=10.5, fontweight="bold")
    fig.tight_layout()
    for suffix, dpi in (("pdf", None), ("svg", None), ("png", 900)):
        target = output / f"figure3_wilke2024_threshold_sensitivity.{suffix}"
        save_options = {"bbox_inches": "tight", "facecolor": "white"}
        if dpi is not None:
            save_options["dpi"] = dpi
        fig.savefig(target, **save_options)
    plt.close(fig)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--figure-dir",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "figures",
    )
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    args.figure_dir.mkdir(parents=True, exist_ok=True)

    source_audit = {name: audit_xlsx(path) for name, path in SOURCES.items() if path.suffix.lower() == ".xlsx"}
    readme = SOURCES["readme"]
    if not readme.is_file():
        raise FileNotFoundError(readme)
    source_audit["readme"] = {
        "path": str(readme),
        "bytes": readme.stat().st_size,
        "sha256": sha256(readme),
    }
    invalid_prior_downloads = []
    for path in EXPECTED_HTML_ERROR_PAGES:
        if path.exists():
            try:
                invalid_prior_downloads.append({**audit_xlsx(path), "reason": "not a valid XLSX; excluded from analysis"})
            except (zipfile.BadZipFile, OSError):
                invalid_prior_downloads.append(
                    {
                        "path": str(path),
                        "bytes": path.stat().st_size,
                        "sha256": sha256(path),
                        "valid_xlsx": False,
                        "reason": "not a valid XLSX; excluded from analysis",
                    }
                )

    used = source_audit["black_mass_composition"]
    if not used.get("valid_xlsx"):
        raise ValueError("Composition workbook failed ZIP integrity checks")
    metrics = extract_metrics(SOURCES["black_mass_composition"])
    threshold_rows = build_threshold_rows(metrics)
    write_csv(args.output / "cell_type_mean_composition.csv", metrics)
    write_csv(args.output / "cutoff_tally_sensitivity.csv", threshold_rows)
    build_figure(metrics, threshold_rows, args.figure_dir)

    selected = {
        str(row["Cu_plus_Al_cutoff_wt_pct"]): row["n_cell_types_passing"]
        for row in threshold_rows
        if row["Cu_plus_Al_cutoff_wt_pct"] in (2.0, 4.0, 7.0, 8.0)
    }
    manifest = {
        "project": "Task 34 public OPARA-276 descriptive threshold audit",
        "run_utc": datetime.now(timezone.utc).isoformat(),
        "analysis_status": "completed; descriptive only",
        "article_doi": ARTICLE_DOI,
        "dataset_doi": DATASET_DOI,
        "dataset_license": "CC BY 4.0 as shown by the OPARA record; attribution required",
        "retrieved_date": "2026-09-28",
        "source_files": source_audit,
        "excluded_invalid_prior_downloads": invalid_prior_downloads,
        "extraction": {
            "worksheet": "Element content",
            "cell_type_codes": "row 8, columns D:M",
            "al_mean_wt_pct": "row 10, columns D:M",
            "cu_mean_wt_pct": "row 12, columns D:M",
            "combined_measure": "arithmetic sum of the separately reported mean Al and Cu mass percentages",
            "dMin_dMax": "not used; no replicate-level paired Cu/Al observations are reconstructed",
        },
        "cutoff_design": {
            "thresholds_wt_pct": "2.0 to 8.0 inclusive, step 0.5",
            "selected_cutoff_wt_pct": 4.0,
            "selected_cutoff_passing_types": selected.get("4.0"),
            "key_tallies_n_of_10": selected,
            "unit_of_tally": "ten source-defined cell types, each counted once",
            "not_estimated": [
                "EU or national feed prevalence",
                "independent batch frequency",
                "assay sensitivity or specificity",
                "direct-regeneration compatibility",
                "route-specific product performance",
            ],
            "used_in_network_model": False,
        },
        "software": {"python": platform.python_version(), "openpyxl": __import__("openpyxl").__version__},
        "outputs": [
            "cell_type_mean_composition.csv",
            "cutoff_tally_sensitivity.csv",
            str(args.figure_dir / "figure3_wilke2024_threshold_sensitivity.pdf"),
            str(args.figure_dir / "figure3_wilke2024_threshold_sensitivity.svg"),
            str(args.figure_dir / "figure3_wilke2024_threshold_sensitivity.png"),
        ],
    }
    manifest["script_sha256"] = sha256(Path(__file__))
    (args.output / "run_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"output": str(args.output), "selected_cutoff_tallies": selected,
                      "source_validity": {k: v.get("valid_xlsx") for k, v in source_audit.items() if k != "readme"}},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
